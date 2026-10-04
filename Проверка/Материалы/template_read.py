from docx import Document
p=r'C:\Users\User\.codex\skills\gost-rpz\assets\Титульный_лист_и_ТЗ.docx'
d=Document(p)
for i,t in enumerate(d.paragraphs): print('P',i,repr(t.text))
for i,t in enumerate(d.tables):
 print('TABLE',i)
 for row in t.rows:print([c.text for c in row.cells])
