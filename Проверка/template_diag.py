from docx import Document
from pathlib import Path
p=Path(r'C:\Users\User\.codex\skills\gost-rpz\assets\Титульный_лист_и_ТЗ.docx')
d=Document(p)
print('sections',len(d.sections))
for i,p in enumerate(d.paragraphs): print('P',i,repr(p.text))
for i,t in enumerate(d.tables):
 print('TABLE',i,len(t.rows),len(t.columns))
 for j,r in enumerate(t.rows):print(j,[c.text for c in r.cells])
