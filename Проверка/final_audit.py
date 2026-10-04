from pathlib import Path
import re,json,zipfile,hashlib
from docx import Document
from docx.oxml.ns import qn
import pdfplumber
R=Path(r'C:\Users\User\Documents\Индивидуальное задание')
M=R/'Отчёты/РПЗ_Последовательная_коррекция.md'
D=R/'Отчёты/РПЗ_Последовательная_коррекция.docx'
d=Document(D);md=M.read_text(encoding='utf8')
imgs=re.findall(r'!\[[^\]]*\]\(([^)]+)\)',md)
assert all((M.parent/x).exists() for x in imgs)
assert md.count('$')%2==0
assert r'\(' not in md and r'\[' not in md
for s in d.sections:
 assert abs(s.page_width.cm-21)<.02 and abs(s.page_height.cm-29.7)<.02
 assert all(abs(a.cm-b)<.02 for a,b in [(s.left_margin,3),(s.right_margin,1),(s.top_margin,2),(s.bottom_margin,2)])
st=d.styles['Normal']
assert st.font.name=='Times New Roman' and st.font.size.pt==14
assert abs(st.paragraph_format.first_line_indent.cm-1.25)<.02
assert st.paragraph_format.line_spacing==1.5
assert len(d.element.findall('.//'+qn('m:oMath')))==112
for t in d.tables[11:]:
 for row in t.rows:
  assert row.height and row.height.cm>=.799
  for c in row.cells:
   for p in c.paragraphs:
    assert all(r.font.size is None or r.font.size.pt==12 for r in p.runs)
caption_style=d.styles['Caption']
while caption_style is not None:
 assert caption_style.font.italic is not True
 caption_style=caption_style.base_style
text='\n'.join(p.text for p in d.paragraphs)
assert all('Рисунок '+str(n) in text for n in range(1,10))
for n in range(1,7):
 assert re.search(r'\['+str(n)+r'\]',md)
with zipfile.ZipFile(D) as z:
 xml=z.read('word/document.xml').decode()
 assert re.search(r'\bTOC\b',xml)
 assert 'Содержание обновляется в Word.' not in xml
assert not any('Error!' in p.text or 'Ошибка! Источник ссылки' in p.text for p in d.paragraphs)
with pdfplumber.open(R/'Рабочие файлы/Проверка/Рендер/release_verified.pdf') as p:
 assert len(p.pages)==33
 pdftexts=[x.extract_text() or '' for x in p.pages]
 assert 'СОДЕРЖАНИЕ' in pdftexts[2]
 assert '1.3 Устойчивость' in pdftexts[8]
 assert 'ПРИЛОЖЕНИЕ А' in pdftexts[32]
 caps={n:next(i+1 for i,t in enumerate(pdftexts) if 'Рисунок '+str(n)+' ' in t) for n in range(1,10)}
 for n,page in caps.items():
  assert len(p.pages[page-1].images)>0
 geometry=json.loads((R/'Рабочие файлы/Проверка/geometry_check.json').read_text(encoding='utf8'))
 assert geometry=={'frequency':[],'dynamic':[]}
checks=re.findall(r'^- \[ \] ([A-Z]\d+) (.+)$',Path(r'C:\Users\User\.codex\skills\gost-rpz\references\checklist.md').read_text(encoding='utf8'),re.M)
unknown={'K01':'Не предоставлены курс, факультет и кафедра; вид — семинарское задание.',
'K03':'Тема едина и содержит 4 слова; дисциплина и группа оставлены для заполнения.',
'K04':'Компоновка, известные ФИО и год проверены; отсутствуют факультет, кафедра, группа, руководитель и даты подписания.',
'K05':'ТЗ включено; полные ФИО и группа не предоставлены. Поле утверждения пустое.',
'K14':'Заполнители реквизитов сохранены намеренно; перед сдачей нужны фактические данные.'}
na={**{k:'Семинарское задание; правила ВКР или практики не применяются.' for k in ['K08','K10','K11','K12','K13','S02','S03']},
'K07':'Работа является семинарским заданием; календарный график курсового проекта к ней не переносится.',
'K09':'Сведения о консультантах не предоставлены; ВКР не выполняется.',
'P03':'Все объекты помещаются на книжных листах.',
'R07':'Альбомных рисунков нет.','R08':'Для рисунков не использованы служебные таблицы.',
'T07':'Все шесть таблиц основной части помещены целиком на одной странице.',
'T09':'Широких альбомных таблиц нет.',
'L01':'В записке нет листингов; код хранится в рабочих файлах.',
'L02':'Листингов нет.','L03':'Листингов нет.','L04':'Листингов нет.',
'B06':'Источников с пятью и более авторами нет.',
'B07':'DOI у использованных учебных файлов и справочных страниц не указан.',
'V06':'Сдаётся электронная записка; печать, подшивка, доклад и утверждение не входят в запрос.'}
notes={
'K02':'Первые страницы — адаптированный официальный титульный лист и ТЗ.',
'K06':'Подписи, утверждения и даты не имитированы; неизвестные поля обозначены.',
'S04':'Содержание Word обновлено целиком; страницы проверены в окончательном PDF.',
'S05':'Кириллические сокращения расположены перед латинскими и греческими символами.',
'S06':'Введение содержит метод, основание, цель и задачи; объём сокращён для семинарского руководства без искусственного наполнения.',
'S08':'Четыре абзаца заключения соотносятся с расчётом, моделями и ограничениями.',
'S09':'Приложение А после источников; ссылка — в пункте 3.4.',
'F01':'Развёрнутые расчётные выражения вынесены на отдельные строки; короткие обозначения и значения используются в тексте.',
'F03':'Формулы не нумерованы; ссылок по номеру формулы нет.',
'F04':'112 нативных редактируемых формул Word, основной размер 14 пт, без использования независимых счётчиков MathType для нумерации формул.',
'B01':'Три предоставленных файла и три официальные справочные страницы SimInTech; метаданные не выдуманы.',
'B02':'Порядок первого цитирования; ссылки 1–6 сопоставлены с записями.',
'V01':'Все поля обновлены Word; содержание и навигация проверены. Нумерация формул не используется.',
'V03':'Просмотрены все страницы окончательного рендера. Страницы 1, 2 и 4 совпадают побайтно с уже просмотренными PNG; остальные 30 просмотрены повторно.',
'V04':'Исправлены перенос заголовка, подписи составных графиков, список и запись источника; повторный рендер — 33 страницы.',
'V05':'Недостающие реквизиты отдельно перечислены; соответствие этих данных не объявлено проверенным.'
}
result=[]
for code,desc in checks:
 status='не проверено' if code in unknown else 'не применимо' if code in na else 'выполнено'
 result.append({'пункт':code,'условие':desc,'результат':status,'основание':unknown.get(code,na.get(code,notes.get(code,'Проверено по свойствам DOCX, тексту и окончательному постраничному рендеру.')))})
out={'страниц':33,'рисунков':9,'таблиц_основной_части':6,'редактируемых_формул':112,'страницы_рисунков':caps,
'математические_границы':'Линейная активная коррекция; рекомендация +40 дБ/дек не выполнена, фактический наклон +60 дБ/дек явно обсуждён.',
'проверка_соединений':geometry,'недостающие_данные':unknown,'контрольный_перечень':result,
'файлы':{x.name:{'байт':x.stat().st_size,'sha256':hashlib.sha256(x.read_bytes()).hexdigest()} for x in [M,D]}}
(R/'Рабочие файлы/Проверка/Итоговая_проверка.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in out.items() if k not in ['контрольный_перечень','недостающие_данные']},ensure_ascii=False,indent=2))
