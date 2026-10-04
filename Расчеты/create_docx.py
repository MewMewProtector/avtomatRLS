from pathlib import Path
import sys,re,copy
R=Path(r'C:\Users\User\Documents\Индивидуальное задание')
sys.path.insert(0,str(R/'Рабочие файлы/Расчёт/packages'))
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
from PIL import Image
mdfile=R/'Отчёты/РПЗ_Последовательная_коррекция.md'
md=mdfile.read_text(encoding='utf8')
# Shorten long sums through defined auxiliary quantities without changing the model.
old=r'$L_{0,\mathrm{ас}}(\omega)=20\log_{10}(5/\omega)+20q_2(\omega)-20q_{2{,}5}(\omega)-20q_1(\omega)-20q_{0{,}8}(\omega)-20q_{0{,}5}(\omega).$'
new=r'$Q_0(\omega)=q_{2{,}5}(\omega)+q_1(\omega)+q_{0{,}8}(\omega)+q_{0{,}5}(\omega),$'+'\n\n'+r'$L_{0,\mathrm{ас}}(\omega)=20\log_{10}(5/\omega)+20q_2(\omega)-20Q_0(\omega).$'
md=md.replace(old,new)
old=r'$L_{\mathrm{ж,ас}}(\omega)=20\log_{10}(10{,}05309649/\omega)-20q_{0{,}05}(\omega)-20q_{0{,}015}(\omega)-20q_{0{,}001}(\omega).$'
new=r'$Q_{\mathrm{ж}}(\omega)=q_{0{,}05}(\omega)+q_{0{,}015}(\omega)+q_{0{,}001}(\omega),$'+'\n\n'+r'$L_{\mathrm{ж,ас}}(\omega)=20\log_{10}(10{,}05309649/\omega)-20Q_{\mathrm{ж}}(\omega).$'
md=md.replace(old,new);mdfile.write_text(md,encoding='utf8')
d=Document(r'C:\Users\User\.codex\skills\gost-rpz\assets\Титульный_лист_и_ТЗ.docx')
title='Последовательная коррекция следящей системы'
def settext(p,t,size=None):
 p.text=t
 for r in p.runs:
  r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
  if size:r.font.size=Pt(size)
def cell(c,t,size=12):
 settext(c.paragraphs[0],t,size)
 for p in c.paragraphs[1:]:p._element.getparent().remove(p._element)
 for p in c.paragraphs:
  p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
d.paragraphs[10].text='К СЕМИНАРСКОМУ ЗАДАНИЮ 3–4'
d.paragraphs[11].text='по [название дисциплины]'
d.paragraphs[33].text='Заведующий кафедрой [указать]'
d.paragraphs[38].text='на выполнение семинарского задания 3–4'
d.paragraphs[44].text='Оформление семинарского задания:'
for r in d.paragraphs[44].runs:r.bold=True;r.italic=True
cell(d.tables[1].cell(0,1),'[факультет]')
cell(d.tables[2].cell(0,1),'[кафедра]')
cell(d.tables[3].cell(1,0),title,14)
cell(d.tables[4].cell(0,0),'Студенты')
cell(d.tables[4].cell(0,5),'И.Р. Косолапов\nЕ.А. Ныркова\nИ.А. Тархов',11)
cell(d.tables[4].cell(0,1),'[группа]')
cell(d.tables[5].cell(0,3),'[И.О. Фамилия]')
cell(d.tables[6].cell(0,0),title,14)
cell(d.tables[6].cell(2,0),'Студенты группы [группа]')
cell(d.tables[6].cell(3,0),'И.Р. Косолапов; Е.А. Ныркова; И.А. Тархов',11)
cell(d.tables[6].cell(4,0),'[полные фамилии, имена и отчества]',10)
cell(d.tables[6].cell(7,0),'Срок выполнения: [указать]')
tz=['Исходная система: вариант 9, единичная отрицательная обратная связь.',
'Спроектировать последовательное КУ методом ЛАЧХ.',
'Требования: tp ≤ 0,5 с (±5 %), σ ≤ 15 %.',
'Запасы: по фазе ≥ 30°, по амплитуде 6–20 дБ.',
'Построить точные и асимптотические ЛАЧХ исходной системы,',
'желаемой системы и их разности — характеристики КУ.',
'Проверить устойчивость, переходный процесс и ошибки слежения.',
'Подготовить подписанные схемы SimInTech и графики.',
'ПИД-регулятор и параллельную коррекцию на этом этапе не выполнять.']
for i,t in enumerate(tz):cell(d.tables[7].cell(i,0),t,11)
cell(d.tables[8].cell(1,0),'Графики и схемы SimInTech с подписями блоков.')
cell(d.tables[8].cell(2,0),'Учебное руководство по последовательной коррекции.')
cell(d.tables[10].cell(3,0),'Студенты')
cell(d.tables[10].cell(3,4),'И.Р. Косолапов; Е.А. Ныркова; И.А. Тархов',10)
for ti,ri in [(6,1),(10,2)]:
 tr=d.tables[ti].rows[ri]._tr;tr.getparent().remove(tr)
# Preserve the official cover/TZ structure; do not fabricate dates or signatures.
for p in d.paragraphs:
 for r in p.runs:
  r.font.name='Times New Roman';r.font.color.rgb=RGBColor(0,0,0)
for sec in d.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(3);sec.right_margin=Cm(1)
 sec.header_distance=Cm(1.25);sec.footer_distance=Cm(1.25)
 sec.different_first_page_header_footer=False
 sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
 for f in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer,sec.even_page_header,sec.even_page_footer]:
  for child in list(f._element):f._element.remove(child)
  f.add_paragraph()
p=d.sections[0].footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
p.add_run('2026 г.').font.size=Pt(14)
# Create page 3, retaining the first two official pages.
sec=d.add_section(WD_SECTION_START.NEW_PAGE);sec.different_first_page_header_footer=False;sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
run=p.add_run();run.font.name='Times New Roman';run.font.size=Pt(12)
b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');run._r.append(b)
ins=OxmlElement('w:instrText');ins.text=' PAGE ';run._r.append(ins)
e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');run._r.append(e)
pg=OxmlElement('w:pgNumType');pg.set(qn('w:start'),'3');sec._sectPr.append(pg)
for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption','TOC 1','TOC 2']:
 try:st=d.styles[name]
 except KeyError:st=d.styles.add_style(name,1)
 st.font.name='Times New Roman';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0)
 st.paragraph_format.first_line_indent=Cm(1.25)
 st.paragraph_format.widow_control=True
 if name.startswith('Heading'):
  st.font.bold=True;st.paragraph_format.keep_with_next=True
  st.paragraph_format.space_after=Pt(21)
  if name=='Heading 1':
   st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.page_break_before=True
  else:
   st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;st.paragraph_format.space_before=Pt(21)
 else:st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
st=d.styles['Caption'];st.font.italic=False;st.font.bold=False;st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['TOC 1','TOC 2']:
 d.styles[name].paragraph_format.first_line_indent=Cm(0);d.styles[name].font.bold=False
d.settings.element.append(OxmlElement('w:updateFields'));d.settings.element[-1].set(qn('w:val'),'true')
transform=etree.XSLT(etree.parse(r'C:\Program Files (x86)\Microsoft Office\Office14\MML2OMML.XSL'))
def math(tex):
 m=etree.fromstring(convert(tex).encode());o=transform(m).getroot()
 for r in o.findall('.//'+qn('m:r')):
  rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'28');rp.append(sz);r.insert(0,rp)
 return copy.deepcopy(o)
def inline(p,t):
 # Render each LaTeX span as a native editable equation.
 t=re.sub(r'\*\*([^*]+)\*\*',r'\1',t)
 for i,x in enumerate(re.split(r'\$([^$]+)\$',t)):
  if i%2:p._p.append(math(x))
  elif x:
   r=p.add_run(x);r.font.name='Times New Roman';r.font.size=Pt(14);r.font.color.rgb=RGBColor(0,0,0)
def blank():
 p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
 return p
def heading(t,lev):
 p=d.add_paragraph(t,style='Heading '+str(lev));return p
# Contents begins on page 3.
p=d.add_paragraph('СОДЕРЖАНИЕ');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(21)
for r in p.runs:r.bold=True
p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
i=OxmlElement('w:instrText');i.text=' TOC \\o "1-2" \\h \\z \\u ';r._r.append(i)
s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'separate');r._r.append(s)
r=p.add_run('Содержание обновляется в Word.')
e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r._r.append(e)
# Skip the abbreviated Markdown front matter; official title/TZ are already copied.
lines=md[md.index('# ПЕРЕЧЕНЬ ОБОЗНАЧЕНИЙ'):].splitlines()
i=0
while i<len(lines):
 s=lines[i].strip()
 if not s:i+=1;continue
 if s.startswith('# '):heading(s[2:],1);i+=1;continue
 if s.startswith('## '):heading(s[3:],2);i+=1;continue
 if s.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',s);path=mdfile.parent/m.group(2)
  w,h=Image.open(path).size;cmw=min(17,(20.5 if 'SimInTech' in str(path) else 19)*w/h)
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
  pic=p.add_run().add_picture(str(path),width=Cm(cmw));pic._inline.docPr.set('descr',m.group(1))
  i+=1
  while i<len(lines) and not lines[i].strip():i+=1
  if i<len(lines) and lines[i].startswith('Рисунок'):
   p=d.add_paragraph(lines[i],style='Caption');p.paragraph_format.keep_together=True;i+=1
  blank();continue
 if s.startswith('Таблица '):
  blank();p=d.add_paragraph(s);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
  i+=1
  while i<len(lines) and not lines[i].strip():i+=1
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   rr=[z.strip() for z in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',z) for z in rr):rows.append(rr)
   i+=1
  t=d.add_table(rows=len(rows),cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  tcpr=t._tbl.tblPr
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   ed=OxmlElement('w:'+edge);ed.set(qn('w:val'),'single');ed.set(qn('w:sz'),'4');ed.set(qn('w:color'),'000000');borders.append(ed)
  tcpr.append(borders)
  for rn,(row,rr) in enumerate(zip(t.rows,rows)):
   row.height=Cm(.8);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
   trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
   if rn==0:
    he=OxmlElement('w:tblHeader');trpr.append(he)
   for c,txt in zip(row.cells,rr):
    c.width=Cm(17/len(rr));c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next=rn<len(rows)-1
    inline(p,txt)
    for r in p.runs:r.font.size=Pt(12);r.font.bold=rn==0
    for z in p._p.findall('.//'+qn('w:sz')):z.set(qn('w:val'),'24')
  blank();continue
 if s.startswith('$') and s.endswith('$'):
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
  p.paragraph_format.keep_together=True;p._p.append(math(s[1:-1]));i+=1;continue
 # Join Markdown soft line breaks in ordinary paragraphs.
 para=[s];i+=1
 while i<len(lines) and lines[i].strip() and not re.match(r'(#|!\[|Таблица |\$|\||\d+[.)])',lines[i].strip()):
  para.append(lines[i].strip());i+=1
 p=d.add_paragraph();inline(p,' '.join(para))
 if re.search(r'(:|равен|равна|имеет вид|определения|из звеньев)$',s):
  p.paragraph_format.keep_with_next=True
 if s.startswith('где'):p.paragraph_format.first_line_indent=Cm(0)
 if s=='Файлы и воспроизводимость':
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.first_line_indent=Cm(0)
  p.paragraph_format.keep_with_next=True
  p.paragraph_format.space_after=Pt(21)
  for r in p.runs:r.font.bold=True
 if re.match(r'\d+[.)]',s):
  p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(0)
 if 'https://' in s:
  for r in p.runs:r.text=r.text.replace('/','/\u200b').replace('_','_\u200b')
d.core_properties.title=title
d.core_properties.subject='Учебное руководство к семинарскому заданию 3–4'
d.core_properties.author='Косолапов И.Р.; Ныркова Е.А.; Тархов И.А.'
d.core_properties.keywords='ЛАЧХ; последовательная коррекция; SimInTech'
out=R/'Отчёты/РПЗ_Последовательная_коррекция.docx'
d.save(out)
print(out,'paragraphs',len(d.paragraphs),'tables',len(d.tables),'equations',len(d.element.findall('.//'+qn('m:oMath'))))
